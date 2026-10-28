"""Finite participant inventories from reviewed public metadata, never image copies.

Regenerate with --source-dir pointing to the original, hash-pinned supporting
files. No network or image transfer is performed by this module.
"""
from __future__ import annotations

import argparse
import csv
from contextlib import closing
import hashlib
import io
import json
from pathlib import Path
import re
import sqlite3
import zipfile

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / 'references/reviewed_missing_participant_inventories_v1.json'
INVENTORY = REFERENCE.with_suffix('.csv')
FIELDS = ('dataset_type', 'short_title', 'download_id', 'participant_id',
          'source_file', 'source_locator', 'evidence_method', 'source_collection')
TCGA = re.compile(r'TCGA-[A-Z0-9]{2}-[A-Z0-9]{4}(?=-|\b)')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def csv_rows(path, encoding='utf-8-sig'):
    return list(csv.DictReader(io.StringIO(path.read_bytes().decode(encoding))))


def workbook_ids(path, column):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook['Metadata']
        rows = sheet.iter_rows(values_only=True)
        index = next(rows).index(column)
        return [(str(row[index]).strip(), f'sheet:Metadata;row:{n};column:{column}')
                for n, row in enumerate(rows, 2) if row[index] is not None]
    finally:
        workbook.close()


def recipe_members(recipe, source_dir):
    """Return only identifiers actually present in the reviewed source grain."""
    name = recipe['source_files'][0]
    path = source_dir / name
    kind = recipe['kind']
    if kind == 'digest':
        return workbook_ids(path, recipe['column'])
    if kind == 'csv':
        return [(r[recipe['column']].strip(), f'row:{n};column:{recipe["column"]}')
                for n, r in enumerate(csv_rows(path, recipe['encoding']), 2)]
    if kind == 'referenced_manifest':
        manifest = (source_dir / recipe['source_files'][1]).read_text()
        series = {line.strip() for line in manifest.splitlines()
                  if re.fullmatch(r'[0-2](?:\.\d+)+', line.strip())}
        if not series:
            raise ValueError('Source manifest contains no series identifiers')
        return [(r['PatientID'].strip(),
                 f'row:{n};ReferencedSeriesInstanceUID:{r["ReferencedSeriesInstanceUID"].strip()}')
                for n, r in enumerate(csv_rows(path, 'latin1'), 2)
                if r['ReferencedSeriesInstanceUID'].strip() in series]
    if kind == 'tcga_workbook':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return [(str(r[0]), f'sheet:CombinedRadiologyPredictors;row:{n};column:Case')
                    for n, r in enumerate(workbook['CombinedRadiologyPredictors'].values, 1)
                    if r and TCGA.fullmatch(str(r[0]))]
        finally:
            workbook.close()
    if kind == 'sdtm_zip':
        with zipfile.ZipFile(path) as archive:
            rows = csv.DictReader(io.StringIO(archive.read('dm.csv').decode('utf-8')))
            return [(r['USUBJID'], f'member:dm.csv;row:{n};STUDYID:{r["STUDYID"]}')
                    for n, r in enumerate(rows, 2)]
    if kind == 'midi_mapping':
        rows = csv_rows(path)
        old, new = [r['id_old'] for r in rows], [r['id_new'] for r in rows]
        if len(set(old)) != len(rows) or len(set(new)) != len(rows):
            raise ValueError('MIDI mapping is not one-to-one')
        curated = {r[0] for r in workbook_ids(source_dir / recipe['source_files'][1], 'Patient ID')}
        if set(new) != curated:
            raise ValueError('MIDI mapping does not match the curated image metadata')
        return [(r['id_new'], f'row:{n};id_old:{r["id_old"]};column:id_new')
                for n, r in enumerate(rows, 2)]
    if kind == 'tcga_text':
        return [(m[0], f'line:{n};TCGA_case_barcode')
                for n, line in enumerate(path.read_text().splitlines(), 1)
                for m in TCGA.finditer(line)]
    if kind == 'til_browser':
        grid = json.loads(path.read_text())['imageGrid']
        if grid.get('paginate') is not False:
            raise ValueError('TIL browser returned a paginated index')
        return [(r['PatientID'], f'imageGrid.values:{n};SlideID:{r["SlideID"]}')
                for n, r in enumerate(grid['values'])]
    if kind == 'til_zip_directory':
        data = json.loads(path.read_text())
        if data.get('range_only') is not True:
            raise ValueError('TIL archive inventory lacks directory-only receipt')
        variants = recipe.get('archive_variants', [recipe.get('archive_variant')])
        if not all(variants) or len(set(variants)) != len(variants):
            raise ValueError('TIL archive folder selection is invalid')
        excluded = recipe.get('excluded_participants', {})
        if excluded and (not recipe.get('idc_presence_review')
                         or set(excluded) & set(recipe['tcga_projects'])):
            raise ValueError('TIL participant exclusions lack a disjoint reviewed decision')
        members = []
        observed = set()
        for variant in variants:
            prefix = variant + '/'
            files = [r for r in data['files'] if r['name'].startswith(prefix)
                     and not r['name'].endswith('/')]
            expected_count = (recipe['map_file_counts'][variant] if 'archive_variants' in recipe
                              else recipe['map_file_count'])
            if len(files) != expected_count:
                raise ValueError('TIL selected variant file count mismatch')
            for entry in files:
                match = re.fullmatch(re.escape(prefix) + r'([a-z]+)/'
                                     r'(TCGA-[A-Z0-9]{2}-[A-Z0-9]{4})-[^/]+\.png', entry['name'])
                if not match or (recipe['tcga_projects'].get(match[2]) or excluded.get(match[2])) != 'TCGA-' + match[1].upper():
                    raise ValueError('TIL archive case/project mapping mismatch')
                observed.add(match[2])
                if match[2] not in excluded:
                    members.append((match[2], 'zip-entry:' + entry['name']))
        if observed != set(recipe['tcga_projects']) | set(excluded):
            raise ValueError('TIL archive selection does not match reviewed participant coverage')
        return members
    raise ValueError(f'Unknown reviewed recipe: {kind}')


def regenerate(source_dir, *, reference=REFERENCE, inventory=INVENTORY):
    metadata = json.loads(reference.read_text())
    for name, source in metadata['source_files'].items():
        if digest(source_dir / name) != source['sha256']:
            raise ValueError(f'Reviewed source digest mismatch: {name}')
    rows = []
    for recipe in metadata['recipes']:
        seen = set()
        for participant, locator in recipe_members(recipe, source_dir):
            participant = participant.strip()
            if not participant or participant in seen:
                continue
            seen.add(participant)
            rows.append(dict(zip(FIELDS, (recipe['dataset_type'], recipe['short_title'],
                recipe['download_id'], participant, recipe['source_files'][0], locator,
                recipe['kind'], recipe.get('source_collection', '')))))
        recipe['participant_count'] = len(seen)
    rows.sort(key=lambda r: tuple(r[k] for k in FIELDS))
    with inventory.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)
    counts = {}
    for row in rows:
        counts.setdefault(row['short_title'], set()).add(row['participant_id'])
    metadata['inventory_sha256'] = digest(inventory)
    metadata['membership_rows'] = len(rows)
    metadata['participants_by_dataset'] = {k: len(v) for k, v in sorted(counts.items())}
    reference.write_text(json.dumps(metadata, indent=2) + '\n')
    return metadata['participants_by_dataset']


def load(reference=REFERENCE, inventory=INVENTORY):
    metadata = json.loads(reference.read_text())
    if digest(inventory) != metadata['inventory_sha256']:
        raise ValueError('Reviewed missing-inventory digest mismatch')
    rows = csv_rows(inventory)
    if len(rows) != metadata['membership_rows']:
        raise ValueError('Reviewed missing-inventory row-count mismatch')
    expected = {(r['dataset_type'], r['short_title'], r['download_id']): r
                for r in metadata['recipes']}
    if len(expected) != len(metadata['recipes']):
        raise ValueError('Duplicate reviewed source recipe')
    for source in metadata['source_files'].values():
        if not re.fullmatch(r'[0-9a-f]{64}', source.get('sha256', '')):
            raise ValueError('Reviewed source lacks a SHA-256 digest')
    actual = {}
    for row in rows:
        key = (row['dataset_type'], row['short_title'], row['download_id'])
        recipe = expected.get(key)
        if (not recipe or row['source_file'] != recipe['source_files'][0]
                or row['evidence_method'] != recipe['kind']
                or row['source_collection'] != recipe.get('source_collection', '')
                or not row['participant_id'].strip() or not row['source_locator']):
            raise ValueError('Reviewed participant row does not match its source recipe')
        identity = (*key, row['participant_id'].casefold())
        if identity in actual:
            raise ValueError('Duplicate reviewed participant row')
        actual[identity] = row
    for key, recipe in expected.items():
        if sum(1 for identity in actual if identity[:3] == key) != recipe['participant_count']:
            raise ValueError('Reviewed recipe participant count mismatch')
        if recipe.get('tcga_projects'):
            identifiers = {r['participant_id'] for identity, r in actual.items() if identity[:3] == key}
            if (set(recipe['tcga_projects']) - identifiers
                    or set(recipe.get('tcga_current_source_ids', [])) - set(recipe['tcga_projects'])
                    or any(not TCGA.fullmatch(pid) or not re.fullmatch(r'TCGA-[A-Z0-9]+', project)
                           for pid, project in recipe['tcga_projects'].items())
                    or not metadata.get('tcga_source_evidence')):
                raise ValueError('TCGA source mapping lacks finite evidence coverage')
        if recipe.get('excluded_participants'):
            included = set(recipe.get('tcga_projects', {}))
            if (not recipe.get('idc_presence_review')
                    or included & set(recipe['excluded_participants'])
                    or set(recipe.get('tcga_current_source_ids', [])) != included):
                raise ValueError('TIL participant exclusions lack complete reviewed IDC coverage')
            review = recipe['idc_presence_review']
            evidence = ROOT / review['evidence_file']
            if (digest(evidence) != review['evidence_sha256']
                    or review['included_cases'] != len(included)
                    or review['excluded_cases'] != len(recipe['excluded_participants'])
                    or review['box_unique_cases'] != len(included) + len(recipe['excluded_participants'])):
                raise ValueError('TIL reviewed IDC evidence digest or coverage mismatch')
            receipt = json.loads(evidence.read_text())
            comparison = next(row for row in receipt['comparison']['rows']
                              if row['source'] == 'box_all_maps')
            if (set(comparison['no_imaging']) != set(recipe['excluded_participants'])
                    or comparison['patients'] - comparison['neither'] != len(included)
                    or receipt['missing_all_idc_collections']['rows']):
                raise ValueError('TIL exclusions differ from the exact reviewed IDC query')
        if recipe.get('participant_species'):
            identifiers = {r['participant_id'] for identity, r in actual.items() if identity[:3] == key}
            if (set(recipe['participant_species']) != identifiers
                    or not recipe.get('classification_review')
                    or set(recipe['participant_species'].values()) - {'Computational Phantom', 'Physical Phantom'}):
                raise ValueError('Reviewed participant classifications lack finite source coverage')
    return metadata, rows


def applicable(metadata, snapshot_db):
    """Bind finite evidence to current visible WordPress download identities."""
    with closing(sqlite3.connect(f'file:{snapshot_db}?mode=ro', uri=True)) as source:
        source.row_factory = sqlite3.Row
        if not source.execute(
            "SELECT 1 FROM sqlite_master WHERE name='agent_current_downloads'"
        ).fetchone():
            # Minimal/offline snapshots with no download catalog cannot activate
            # any of these source-bound inventories.
            return {}
        active = {(r['dataset_type'], r['short_title'], str(r['download_id'])): dict(r)
                  for r in source.execute('SELECT * FROM agent_current_downloads WHERE hidden=0')}
    recipes = {}
    for recipe in metadata['recipes']:
        key = (recipe['dataset_type'], recipe['short_title'], recipe['download_id'])
        if key not in active:
            continue
        for binding in recipe['bindings']:
            matches = [r for k, r in active.items()
                       if k[1:] == (binding['short_title'], binding['download_id'])]
            if len(matches) != 1 or matches[0].get(binding['field']) != binding['url']:
                raise ValueError(f'Current download differs from reviewed evidence: {key}')
        recipes[key] = (recipe, active[key])
    return recipes


def ingest(conn, snapshot_db, *, api, reference=REFERENCE, inventory=INVENTORY):
    metadata, rows = load(reference, inventory)
    recipes = applicable(metadata, snapshot_db)
    count = 0
    for row in rows:
        scope = (row['dataset_type'], row['short_title'], row['download_id'])
        if scope not in recipes:
            continue
        recipe, download = recipes[scope]
        source = metadata['source_files'][row['source_file']]
        evidence = {**row, 'source_url': source['url'], 'source_sha256': source['sha256'],
                    'source_files': {n: metadata['source_files'][n] for n in recipe['source_files']},
                    'reviewed_at': metadata['reviewed_at'], 'metadata_only': True}
        if recipe.get('classification_review'):
            evidence['classification_review'] = recipe['classification_review']
        if recipe.get('idc_presence_review'):
            evidence['idc_presence_review'] = recipe['idc_presence_review']
        key = api.ensure_participant(conn, row['dataset_type'], row['short_title'],
            row['participant_id'], managed_system='tcia_wordpress',
            namespace=f'tcia_dataset:{row["short_title"]}',
            evidence='official_reviewed_participant_metadata', provenance=evidence)
        # This points to public metadata documenting membership. It is neither an
        # image file nor a claim that a participant has locally inventoried images.
        api.add_asset(conn, {
            'participant_asset_id': api.stable_id('pa', key, 'reviewed_participant_metadata', row['download_id']),
            'participant_key': key, 'managed_system': 'tcia_wordpress',
            'source_artifact': 'reviewed_missing_participant_inventories_v1',
            'access_level': download['access_level'], 'data_domain': 'participant_metadata',
            'media_kind': 'tabular', 'modality': '', 'file_format': '',
            'object_role': 'metadata_only', 'study_count': None, 'series_count': None,
            'file_count': None, 'known_size_bytes': None, 'has_file_level_metadata': 0,
            'detail_pointer': source['url'], 'access_route': '',
            'inventory_status': 'known_metadata_authorization_required' if download['controlled_access'] else 'known',
            'source_version': metadata['reviewed_at'], 'provenance_json': api.json_dumps(evidence),
        })
        count += 1
    api.record_source(conn, 'reviewed_missing_participant_inventories', inventory, True,
                      count, 'Hash-pinned public metadata establishes participant membership; no image/file counts are inferred.')
    return count


def link_sources(conn, snapshot_db, *, api, reference=REFERENCE, inventory=INVENTORY):
    metadata, rows = load(reference, inventory)
    recipes = applicable(metadata, snapshot_db)
    parents_by_id = {}
    for row in rows:
        if row['dataset_type'] == 'Collection' and row['evidence_method'] in {'referenced_manifest', 'digest'}:
            parents_by_id.setdefault((row['short_title'], row['participant_id']), []).append(row)
    count = 0
    for row in rows:
        scope = (row['dataset_type'], row['short_title'], row['download_id'])
        parent = row['source_collection']
        if not parent or scope not in recipes:
            continue
        # Require a separately verified parent row: a referenced-series match to
        # its manifest, or its own patient-level NBIA metadata spreadsheet.
        parents = parents_by_id.get((parent, row['participant_id']), [])
        if len(parents) != 1:
            raise ValueError('Annotation participant lacks verified parent metadata membership')
        parent_row = parents[0]
        if (parent_row['dataset_type'], parent, parent_row['download_id']) not in recipes:
            continue
        target = api.participant_key('Collection', parent, row['participant_id'])
        if not conn.execute('SELECT 1 FROM participants WHERE participant_key=?', (target,)).fetchone():
            raise ValueError('Reviewed source participant is absent')
        api.record_source_collection_link(conn,
            analysis_result_participant_key=api.participant_key(row['dataset_type'], row['short_title'], row['participant_id']),
            source_collection_short_title=parent, source_participant_id=row['participant_id'],
            source_collection_participant_key=target, link_status='linked_source_collection',
            resolution_method='official_reviewed_finite_participant_crosswalk', confidence='high',
            evidence_source=reference.name, evidence_pointer=parent_row['source_locator'],
            evidence_url=metadata['source_files'][row['source_file']]['url'],
            evidence={'source_row': parent_row, 'source_files': {
                name: metadata['source_files'][name] for name in
                recipes[(parent_row['dataset_type'], parent, parent_row['download_id'])][0]['source_files']},
                      'reviewed_at': metadata['reviewed_at']})
        count += 1
    count += link_tcga_sources(conn, metadata, rows, recipes, api=api, reference=reference)
    return count


def link_tcga_sources(conn, metadata, rows, recipes, *, api, reference):
    """Preserve original TCGA project/case identity, including external sources.

    No Collection rows are manufactured when a GDC project is not TCIA-published.
    Current-source confirmation comes from the pinned original-SM IDC projection,
    separately from membership in an existing TCIA Collection.
    """
    count = 0
    for row in rows:
        scope = (row['dataset_type'], row['short_title'], row['download_id'])
        if scope not in recipes:
            continue
        recipe = recipes[scope][0]
        project = recipe.get('tcga_projects', {}).get(row['participant_id'])
        if not project:
            continue
        left = api.participant_key(*scope[:2], row['participant_id'])
        target = api.participant_key('Collection', project, row['participant_id'])
        present = conn.execute('SELECT 1 FROM participants WHERE participant_key=?', (target,)).fetchone()
        current = row['participant_id'] in recipe['tcga_current_source_ids']
        api.record_source_collection_link(conn,
            analysis_result_participant_key=left, source_collection_short_title=project,
            source_participant_id=row['participant_id'], source_collection_participant_key=target if present else None,
            link_status='linked_source_collection' if present else 'source_participant_not_current',
            resolution_method='official_reviewed_finite_participant_crosswalk',
            confidence='high', evidence_source=reference.name,
            evidence_pointer=row['source_locator'],
            evidence_url=metadata['source_files'][row['source_file']]['url'],
            evidence={'tcga_original_source': True, 'project': project, 'case_id': row['participant_id'],
                      'current_original_sm_confirmed': current,
                      'source_sha256': metadata['source_files'][row['source_file']]['sha256'],
                      'tcga_source_evidence': metadata['tcga_source_evidence'],
                      'source_locator': row['source_locator']})
        count += 1
    return count


def tcga_source_groups(conn, *, api, reference=REFERENCE, inventory=INVENTORY):
    """Finite TCGA source anchors for deduplication across derived datasets."""
    metadata, rows = load(reference, inventory)
    allowed = {}
    for recipe in metadata['recipes']:
        for identifier, project in recipe.get('tcga_projects', {}).items():
            # TCGA's documented participant barcode identifies the original
            # participant independently of current source-image availability.
            allowed[(api.participant_key(recipe['dataset_type'], recipe['short_title'], identifier),
                     project, identifier)] = True
    groups = {}
    for row in conn.execute('SELECT * FROM participant_source_links WHERE evidence_source=? AND confidence=\'high\'',
                            (reference.name,)):
        key = (row['analysis_result_participant_key'], row['source_collection_short_title'], row['source_participant_id'])
        if key in allowed:
            groups.setdefault(('GDC_TCGA', key[1], key[2]), []).append(key[0])
    return groups


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(regenerate(args.source_dir), indent=2))
